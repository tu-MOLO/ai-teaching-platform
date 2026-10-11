import sys
import types
from datetime import datetime
from types import SimpleNamespace

import pytest

from app.services.export import ExportService


def _make_lesson_plan(**overrides) -> SimpleNamespace:
    defaults = dict(
        title="测试教案",
        subject="数学",
        grade="三年级",
        duration=45,
        status=SimpleNamespace(value="draft"),
        teaching_goals_a="基础目标",
        teaching_goals_b="提高目标",
        teaching_goals_c="拓展目标",
        teaching_objectives="总体教学目标",
        teaching_content="教学内容文本",
        teaching_methods="教学方法文本",
        teaching_process="教学过程文本",
        teaching_resources="教学资源文本",
        assessment="评价方式文本",
        notes="备注文本",
    )
    defaults.update(overrides)
    return SimpleNamespace(**defaults)


def _make_student(**overrides) -> SimpleNamespace:
    defaults = dict(
        name="张三",
        gender="male",
        birth_date="2015-05-10",
        grade="三年级",
        class_name="1班",
        parent_contact="13800138000",
    )
    defaults.update(overrides)
    return SimpleNamespace(**defaults)


def _make_portfolio(**overrides) -> SimpleNamespace:
    defaults = dict(
        type="work",
        title="作品一",
        content="作品内容",
        created_at=datetime(2024, 6, 15, 10, 30, 0),
        cognitive_score=None,
        skill_score=None,
        creativity_score=None,
        cooperation_score=None,
        attention_score=None,
    )
    defaults.update(overrides)
    return SimpleNamespace(**defaults)


class TestExportToWord:
    def test_generates_valid_docx_bytes(self):
        service = ExportService()
        lp = _make_lesson_plan()
        result = service.export_to_word(lp)
        assert isinstance(result, bytes)
        assert len(result) > 0
        assert result[:2] == b"PK"

    def test_docx_contains_title(self):
        from io import BytesIO

        from docx import Document

        service = ExportService()
        lp = _make_lesson_plan(title="特殊教案标题")
        result = service.export_to_word(lp)
        doc = Document(BytesIO(result))
        assert doc.paragraphs[0].text == "特殊教案标题"

    def test_docx_with_none_optional_fields(self):
        service = ExportService()
        lp = _make_lesson_plan(
            teaching_content=None,
            teaching_methods=None,
            teaching_process=None,
            teaching_resources=None,
            assessment=None,
            notes=None,
        )
        result = service.export_to_word(lp)
        assert isinstance(result, bytes)
        assert len(result) > 0


class TestGenerateHtml:
    def test_generates_html_with_all_sections(self):
        service = ExportService()
        lp = _make_lesson_plan()
        html = service._generate_html(lp)
        assert "<!DOCTYPE html>" in html
        assert "测试教案" in html
        assert "数学" in html
        assert "三年级" in html
        assert "45分钟" in html
        assert "draft" in html
        assert "A层（基础）" in html
        assert "B层（提高）" in html
        assert "C层（拓展）" in html
        assert "教学内容" in html
        assert "教学方法" in html
        assert "教学过程" in html
        assert "教学资源" in html
        assert "评价方式" in html
        assert "备注" in html

    def test_handles_none_fields(self):
        service = ExportService()
        lp = _make_lesson_plan(
            teaching_goals_a=None,
            teaching_goals_b=None,
            teaching_goals_c=None,
            teaching_content=None,
            teaching_methods=None,
            teaching_process=None,
            teaching_resources=None,
            assessment=None,
            notes=None,
        )
        html = service._generate_html(lp)
        assert "A层" not in html
        assert "B层" not in html
        assert "C层" not in html
        assert "教学内容" not in html
        assert "教学方法" not in html
        assert "教学过程" not in html

    def test_html_escapes_special_chars(self):
        service = ExportService()
        lp = _make_lesson_plan(title="<script>alert('xss')</script>")
        html = service._generate_html(lp)
        assert "<script>" not in html
        assert "&lt;script&gt;" in html

    def test_html_contains_status_value(self):
        service = ExportService()
        lp = _make_lesson_plan(status=SimpleNamespace(value="published"))
        html = service._generate_html(lp)
        assert "published" in html


class TestBuildLessonPlanLines:
    def test_generates_text_lines(self):
        service = ExportService()
        lp = _make_lesson_plan()
        lines = service._build_lesson_plan_lines(lp)
        assert isinstance(lines, list)
        assert any("测试教案" in line for line in lines)
        assert any("数学" in line for line in lines)
        assert any("三年级" in line for line in lines)
        assert any("45" in line for line in lines)
        assert any("教学目标" in line for line in lines)
        assert any("教学内容" in line for line in lines)
        assert any("教学方法" in line for line in lines)
        assert any("教学过程" in line for line in lines)
        assert any("教学资源" in line for line in lines)
        assert any("备注" in line for line in lines)

    def test_none_fields_show_default(self):
        service = ExportService()
        lp = _make_lesson_plan(
            teaching_objectives=None,
            teaching_content=None,
            teaching_methods=None,
        )
        lines = service._build_lesson_plan_lines(lp)
        assert any("无" in line for line in lines)

    def test_status_uses_value_attribute(self):
        service = ExportService()
        lp = _make_lesson_plan(status=SimpleNamespace(value="published"))
        lines = service._build_lesson_plan_lines(lp)
        assert any("published" in line for line in lines)


class TestBuildPortfolioReportLines:
    def test_generates_text_lines_for_portfolio(self):
        service = ExportService()
        student = _make_student()
        portfolio = _make_portfolio()
        lines = service._build_portfolio_report_lines(student, [portfolio])
        assert any("张三" in line for line in lines)
        assert any("作品一" in line for line in lines)
        assert any("作品" in line for line in lines)

    def test_empty_portfolios(self):
        service = ExportService()
        student = _make_student()
        lines = service._build_portfolio_report_lines(student, [])
        assert any("暂无成长记录" in line for line in lines)

    def test_includes_scores(self):
        service = ExportService()
        student = _make_student()
        portfolio = _make_portfolio(
            cognitive_score=85,
            skill_score=90,
        )
        lines = service._build_portfolio_report_lines(student, [portfolio])
        joined = " ".join(lines)
        assert "认知" in joined
        assert "技能" in joined

    def test_student_info_in_lines(self):
        service = ExportService()
        student = _make_student()
        lines = service._build_portfolio_report_lines(student, [])
        joined = " ".join(lines)
        assert "张三" in joined
        assert "male" in joined
        assert "三年级" in joined
        assert "1班" in joined


class TestGeneratePortfolioHtml:
    def test_generates_html_with_student_info(self):
        service = ExportService()
        student = _make_student()
        html = service._generate_portfolio_html(student, [])
        assert "张三" in html
        assert "male" in html
        assert "三年级" in html
        assert "1班" in html
        assert "13800138000" in html

    def test_groups_portfolios_by_type(self):
        service = ExportService()
        student = _make_student()
        portfolios = [
            _make_portfolio(type="work", title="作品A"),
            _make_portfolio(type="evaluation", title="评价B"),
            _make_portfolio(type="work", title="作品C"),
        ]
        html = service._generate_portfolio_html(student, portfolios)
        assert "作品" in html
        assert "评价" in html
        assert "作品A" in html
        assert "评价B" in html
        assert "作品C" in html

    def test_includes_evaluation_scores(self):
        service = ExportService()
        student = _make_student()
        portfolio = _make_portfolio(
            cognitive_score=85,
            skill_score=90,
            creativity_score=78,
            cooperation_score=88,
            attention_score=92,
        )
        html = service._generate_portfolio_html(student, [portfolio])
        assert "多维度评价" in html
        assert "认知" in html
        assert "技能" in html
        assert "创意" in html
        assert "合作" in html
        assert "注意力" in html

    def test_no_scores_no_evaluation_section(self):
        service = ExportService()
        student = _make_student()
        portfolio = _make_portfolio()
        html = service._generate_portfolio_html(student, [portfolio])
        assert "多维度评价" not in html

    def test_content_rendered_when_present(self):
        service = ExportService()
        student = _make_student()
        portfolio = _make_portfolio(content="这是作品内容")
        html = service._generate_portfolio_html(student, [portfolio])
        assert "这是作品内容" in html

    def test_content_absent_no_paragraph(self):
        service = ExportService()
        student = _make_student()
        portfolio = _make_portfolio(content=None)
        html = service._generate_portfolio_html(student, [portfolio])
        assert "<p></p>" not in html.replace("\n", "").replace(" ", "")

    def test_type_name_mapping(self):
        service = ExportService()
        student = _make_student()
        portfolios = [
            _make_portfolio(type="milestone", title="里程碑项"),
            _make_portfolio(type="observation", title="观察项"),
        ]
        html = service._generate_portfolio_html(student, portfolios)
        assert "里程碑" in html
        assert "观察" in html

    def test_unknown_type_uses_raw_key(self):
        service = ExportService()
        student = _make_student()
        portfolio = _make_portfolio(type="custom_type", title="自定义项")
        html = service._generate_portfolio_html(student, [portfolio])
        assert "custom_type" in html

    def test_parent_contact_default(self):
        service = ExportService()
        student = _make_student(parent_contact=None)
        html = service._generate_portfolio_html(student, [])
        assert "未提供" in html


# ============================================================
# WeasyPrint 可用时的 PDF 成功路径（测试中注入桩模块）
# ============================================================


def _install_fake_weasyprint(monkeypatch, pdf: bytes = b"%PDF-FAKE") -> None:
    """向 sys.modules 注入 weasyprint 桩，使函数内延迟导入成功。"""
    module = types.ModuleType("weasyprint")

    class _HTML:
        last_instance: "_HTML | None" = None

        def __init__(self, string: str):
            self.string = string
            _HTML.last_instance = self

        def write_pdf(self, stylesheets=None):
            self.stylesheets = stylesheets
            return pdf

    class _CSS:
        def __init__(self, string: str):
            self.string = string

    setattr(module, "HTML", _HTML)
    setattr(module, "CSS", _CSS)
    monkeypatch.setitem(sys.modules, "weasyprint", module)


def _make_iep(**overrides) -> SimpleNamespace:
    defaults = dict(
        id="iep-1",
        student_name="李四",
        student=None,
        term="2026 春季学期",
        status="locked",
        revision_no=2,
        reviewed_at=datetime(2026, 6, 30, 9, 30),
        student_profile="学生基本情况：喜欢音乐，对声音敏感。",
        evaluation_plan="每周一次自然观察并记录。",
        supports={
            "teaching": ["视觉时间表"],
            "rehab": [],
            "family": ["家中固定作息"],
        },
        domain_assessments=[
            SimpleNamespace(
                domain_name="语言",
                overall_level=None,
                summary="词汇量增长。",
                items=[
                    SimpleNamespace(item_name="听懂一步指令", level=3, note="稳定"),
                    SimpleNamespace(item_name="主动表达需求", level=None, note=None),
                ],
            ),
        ],
        long_term_goals=[
            SimpleNamespace(
                domain_name="语言",
                goal="提升日常沟通能力",
                short_term_goals=[
                    SimpleNamespace(
                        condition=None,
                        behavior="用短句表达需求",
                        criterion=None,
                        support_level=None,
                        status="not_started",
                        teaching_methods=None,
                    ),
                    SimpleNamespace(
                        condition="餐前",
                        behavior="提出请求",
                        criterion="连续3天",
                        support_level=2,
                        status="achieved",
                        teaching_methods=["示范", "强化"],
                    ),
                ],
            ),
        ],
    )
    defaults.update(overrides)
    return SimpleNamespace(**defaults)


class TestWeasyPrintPdfPaths:
    def test_lesson_plan_pdf_uses_weasyprint(self, monkeypatch):
        _install_fake_weasyprint(monkeypatch)
        service = ExportService()
        result = service.export_to_pdf(_make_lesson_plan())
        assert result == b"%PDF-FAKE"

    def test_portfolio_pdf_uses_weasyprint(self, monkeypatch):
        _install_fake_weasyprint(monkeypatch, b"%PDF-PORTFOLIO")
        service = ExportService()
        result = service.export_student_portfolio_to_pdf(_make_student(), [_make_portfolio()])
        assert result == b"%PDF-PORTFOLIO"

    def test_iep_pdf_uses_weasyprint(self, monkeypatch):
        _install_fake_weasyprint(monkeypatch)
        service = ExportService()
        result = service.export_iep_to_pdf(_make_iep())
        assert result[:5] == b"%PDF-"

    def test_word_raises_importerror_when_docx_missing(self, monkeypatch):
        monkeypatch.setitem(sys.modules, "docx", None)
        service = ExportService()
        with pytest.raises(ImportError):
            service.export_to_word(_make_lesson_plan())


class TestIEPHtml:
    def test_full_iep_html_contains_all_sections(self):
        service = ExportService()
        html = service._generate_iep_html(_make_iep())
        assert "个别化教育计划" in html
        assert "李四" in html
        assert "2026 春季学期" in html
        assert "第 2 版" in html
        assert "已定稿" in html
        assert "2026-06-30 09:30" in html
        # 领域与条目（含未评估、空备注回退）
        assert "语言" in html
        assert "听懂一步指令" in html
        assert "3 独立完成" in html
        assert "未评估" in html
        # 目标
        assert "提升日常沟通能力" in html
        assert "用短句表达需求" in html
        assert "未开始" in html
        assert "已达成" in html
        assert "示范、强化" in html
        # 支持措施与评价方式
        assert "视觉时间表" in html
        assert "家中固定作息" in html
        assert "每周一次自然观察并记录" in html

    def test_student_name_falls_back_to_related_student(self):
        service = ExportService()
        iep = _make_iep(
            student_name=None,
            student=SimpleNamespace(name="王小明"),
        )
        assert service._iep_student_name(iep) == "王小明"
        html = service._generate_iep_html(iep)
        assert "王小明" in html

    def test_empty_assessments_and_goals_placeholders(self):
        service = ExportService()
        iep = _make_iep(domain_assessments=[], long_term_goals=[])
        html = service._generate_iep_html(iep)
        assert "暂无评估" in html
        assert "暂无目标" in html


class TestSimplePdfRenderers:
    def test_reportlab_renderer_paginates_long_content(self):
        # A4 每页约 41 行，80 行触发 showPage 分页分支
        service = ExportService()
        result = service._render_simple_pdf_reportlab([f"第{i}行" for i in range(80)])
        assert result[:5] == b"%PDF-"
        assert len(result) > 0

    def test_falls_back_to_legacy_pdf_when_reportlab_fails(self, monkeypatch):
        service = ExportService()

        def _boom(_lines):
            raise RuntimeError("reportlab unavailable")

        monkeypatch.setattr(service, "_render_simple_pdf_reportlab", _boom)

        # 80 行超过 legacy 每页 38 行 -> 3 页
        result = service._render_simple_pdf([f"第{i}行" for i in range(80)])
        assert result[:8] == b"%PDF-1.4"
        assert b"/Count 3" in result
        assert b"startxref" in result

    def test_iep_export_falls_back_to_simple_pdf_without_weasyprint(self, monkeypatch):
        # 模拟 weasyprint 缺失：sys.modules 中放入 None 令导入失败
        monkeypatch.setitem(sys.modules, "weasyprint", None)
        service = ExportService()
        result = service.export_iep_to_pdf(_make_iep())
        assert result[:5] == b"%PDF-"
