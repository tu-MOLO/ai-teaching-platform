"""
导出服务
"""

from html import escape
from io import BytesIO
from typing import Any


class ExportService:
    """
    导出服务类
    """

    def export_to_pdf(self, lesson_plan) -> bytes:
        """
        导出教案为PDF

        Args:
            lesson_plan: 教案对象

        Returns:
            PDF文件内容（字节流）
        """
        try:
            from weasyprint import CSS, HTML
        except ImportError:
            return self._render_simple_pdf(self._build_lesson_plan_lines(lesson_plan))

        # 构建HTML内容
        html_content = self._generate_html(lesson_plan)

        # 生成PDF
        css = CSS(string="""
            body {
                font-family: SimHei, Arial, sans-serif;
                margin: 20px;
            }
            h1 {
                color: #333;
                text-align: center;
                margin-bottom: 30px;
            }
            h2 {
                color: #555;
                margin-top: 20px;
                margin-bottom: 10px;
                border-bottom: 1px solid #ddd;
                padding-bottom: 5px;
            }
            p {
                line-height: 1.5;
                margin-bottom: 10px;
            }
            .section {
                margin-bottom: 20px;
            }
            .goal-level {
                margin-left: 20px;
                margin-bottom: 10px;
            }
            .goal-level h3 {
                color: #666;
                margin-bottom: 5px;
            }
        """)

        pdf_bytes: bytes = HTML(string=html_content).write_pdf(  # type: ignore[no-any-return]
            stylesheets=[css]
        )
        return pdf_bytes

    def _add_teaching_goals_to_doc(self, doc, lesson_plan) -> None:
        """添加分层教学目标到Word文档"""
        doc.add_heading("教学目标", level=1)
        if lesson_plan.teaching_goals_a:
            doc.add_heading("A层（基础）", level=2)
            doc.add_paragraph(escape(lesson_plan.teaching_goals_a or ""))
        if lesson_plan.teaching_goals_b:
            doc.add_heading("B层（提高）", level=2)
            doc.add_paragraph(escape(lesson_plan.teaching_goals_b or ""))
        if lesson_plan.teaching_goals_c:
            doc.add_heading("C层（拓展）", level=2)
            doc.add_paragraph(escape(lesson_plan.teaching_goals_c or ""))

    def _add_optional_section_to_doc(self, doc, heading: str, content: Any) -> None:
        """添加可选章节到Word文档"""
        if content:
            doc.add_heading(heading, level=1)
            doc.add_paragraph(escape(content or ""))

    def export_to_word(self, lesson_plan) -> bytes:
        """
        导出教案为Word

        Args:
            lesson_plan: 教案对象

        Returns:
            Word文件内容（字节流）
        """
        try:
            from docx import Document
        except ImportError:
            raise ImportError("Word export not supported: python-docx dependencies missing")

        doc = Document()

        # 添加标题
        doc.add_heading(lesson_plan.title, 0)

        # 添加基本信息
        doc.add_heading("基本信息", level=1)
        doc.add_paragraph(f'学科: {escape(lesson_plan.subject or "")}')
        doc.add_paragraph(f'年级: {escape(lesson_plan.grade or "")}')
        doc.add_paragraph(f"课时时长: {lesson_plan.duration}分钟")
        doc.add_paragraph(f"状态: {lesson_plan.status.value}")

        self._add_teaching_goals_to_doc(doc, lesson_plan)
        self._add_optional_section_to_doc(doc, "教学内容", lesson_plan.teaching_content)
        self._add_optional_section_to_doc(doc, "教学方法", lesson_plan.teaching_methods)
        self._add_optional_section_to_doc(doc, "教学过程", lesson_plan.teaching_process)
        self._add_optional_section_to_doc(doc, "教学资源", lesson_plan.teaching_resources)
        self._add_optional_section_to_doc(doc, "评价方式", lesson_plan.assessment)
        self._add_optional_section_to_doc(doc, "备注", lesson_plan.notes)

        # 保存为字节流
        stream = BytesIO()
        doc.save(stream)
        stream.seek(0)
        return stream.getvalue()

    def _generate_html(self, lesson_plan) -> str:
        """
        生成HTML内容

        Args:
            lesson_plan: 教案对象

        Returns:
            HTML字符串
        """
        html = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="UTF-8">
            <title>{escape(lesson_plan.title or '')}</title>
        </head>
        <body>
            <h1>{escape(lesson_plan.title or '')}</h1>

            <div class="section">
                <h2>基本信息</h2>
                <p>学科: {escape(lesson_plan.subject or '')}</p>
                <p>年级: {escape(lesson_plan.grade or '')}</p>
                <p>课时时长: {lesson_plan.duration}分钟</p>
                <p>状态: {escape(str(lesson_plan.status.value))}</p>
            </div>

            <div class="section">
                <h2>教学目标</h2>
        """

        # 添加分层教学目标
        if lesson_plan.teaching_goals_a:
            html += f"""
                <div class="goal-level">
                    <h3>A层（基础）</h3>
                    <p>{escape(lesson_plan.teaching_goals_a or '')}</p>
                </div>
            """

        if lesson_plan.teaching_goals_b:
            html += f"""
                <div class="goal-level">
                    <h3>B层（提高）</h3>
                    <p>{escape(lesson_plan.teaching_goals_b or '')}</p>
                </div>
            """

        if lesson_plan.teaching_goals_c:
            html += f"""
                <div class="goal-level">
                    <h3>C层（拓展）</h3>
                    <p>{escape(lesson_plan.teaching_goals_c or '')}</p>
                </div>
            """

        # 添加其他内容
        sections = [
            ("教学内容", lesson_plan.teaching_content),
            ("教学方法", lesson_plan.teaching_methods),
            ("教学过程", lesson_plan.teaching_process),
            ("教学资源", lesson_plan.teaching_resources),
            ("评价方式", lesson_plan.assessment),
            ("备注", lesson_plan.notes),
        ]

        for section_name, content in sections:
            if content:
                html += f"""
            <div class="section">
                <h2>{section_name}</h2>
                <p>{escape(content or '')}</p>
            </div>
        """

        html += """
        </body>
        </html>
        """

        return html

    def export_student_portfolio_to_pdf(self, student, portfolios) -> bytes:
        """
        导出学生成长报告为PDF

        Args:
            student: 学生对象
            portfolios: 学生的成长档案列表

        Returns:
            PDF文件内容（字节流）
        """
        try:
            from weasyprint import CSS, HTML
        except (ImportError, OSError):
            return self._render_simple_pdf(self._build_portfolio_report_lines(student, portfolios))

        # 构建HTML内容
        html_content = self._generate_portfolio_html(student, portfolios)

        # 生成PDF
        css = CSS(string="""
            body {
                font-family: SimHei, Arial, sans-serif;
                margin: 20px;
            }
            h1 {
                color: #333;
                text-align: center;
                margin-bottom: 30px;
            }
            h2 {
                color: #555;
                margin-top: 30px;
                margin-bottom: 15px;
                border-bottom: 1px solid #ddd;
                padding-bottom: 5px;
            }
            h3 {
                color: #666;
                margin-top: 20px;
                margin-bottom: 10px;
            }
            p {
                line-height: 1.5;
                margin-bottom: 10px;
            }
            .section {
                margin-bottom: 20px;
            }
            .student-info {
                background-color: #f5f5f5;
                padding: 15px;
                border-radius: 5px;
                margin-bottom: 30px;
            }
            .portfolio-item {
                border: 1px solid #ddd;
                padding: 15px;
                border-radius: 5px;
                margin-bottom: 15px;
            }
            .portfolio-title {
                font-weight: bold;
                font-size: 16px;
                margin-bottom: 5px;
            }
            .portfolio-meta {
                font-size: 14px;
                color: #666;
                margin-bottom: 10px;
            }
            .evaluation-scores {
                margin-top: 10px;
                padding: 10px;
                background-color: #f9f9f9;
                border-radius: 5px;
            }
            .score-item {
                display: inline-block;
                margin-right: 20px;
                margin-bottom: 5px;
            }
            .score-label {
                font-weight: bold;
                margin-right: 5px;
            }
        """)

        pdf_bytes: bytes = HTML(string=html_content).write_pdf(  # type: ignore[no-any-return]
            stylesheets=[css]
        )
        return pdf_bytes

    def _generate_portfolio_html(self, student, portfolios) -> str:
        """
        生成学生成长报告的HTML内容

        Args:
            student: 学生对象
            portfolios: 学生的成长档案列表

        Returns:
            HTML字符串
        """
        html = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="UTF-8">
            <title>{escape(student.name or '')}的成长报告</title>
        </head>
        <body>
            <h1>{escape(student.name or '')}的成长报告</h1>

            <div class="student-info">
                <h2>学生基本信息</h2>
                <p>姓名: {escape(student.name or '')}</p>
                <p>性别: {escape(student.gender or '')}</p>
                <p>出生日期: {student.birth_date}</p>
                <p>年级: {escape(student.grade or '')}</p>
                <p>班级: {escape(student.class_name or '')}</p>
                <p>家长联系方式: {escape(student.parent_contact or '未提供')}</p>
            </div>
        """

        # 按类型分组成长档案
        portfolios_by_type: dict[str, list[Any]] = {}
        for portfolio in portfolios:
            if portfolio.type not in portfolios_by_type:
                portfolios_by_type[portfolio.type] = []
            portfolios_by_type[portfolio.type].append(portfolio)

        # 类型名称映射
        type_names = {
            "work": "作品",
            "evaluation": "评价",
            "observation": "观察",
            "milestone": "里程碑",
        }

        # 添加各类型的成长档案
        for type_key, type_portfolios in portfolios_by_type.items():
            type_name = type_names.get(type_key, type_key)
            html += f"""
            <div class="section">
                <h2>{escape(type_name)}</h2>
            """

            for portfolio in type_portfolios:
                html += f"""
                <div class="portfolio-item">
                    <div class="portfolio-title">{escape(portfolio.title or '')}</div>
                    <div class="portfolio-meta">创建时间: {portfolio.created_at.strftime('%Y-%m-%d %H:%M:%S')}</div>
                    {f'<p>{escape(portfolio.content or "")}</p>' if portfolio.content else ''}
                """

                # 添加评价分数
                if any(
                    [
                        portfolio.cognitive_score,
                        portfolio.skill_score,
                        portfolio.creativity_score,
                        portfolio.cooperation_score,
                        portfolio.attention_score,
                    ]
                ):
                    html += """
                    <div class="evaluation-scores">
                        <h3>多维度评价</h3>
                    """

                    scores = [
                        ("认知", portfolio.cognitive_score),
                        ("技能", portfolio.skill_score),
                        ("创意", portfolio.creativity_score),
                        ("合作", portfolio.cooperation_score),
                        ("注意力", portfolio.attention_score),
                    ]

                    for score_name, score_value in scores:
                        if score_value is not None:
                            html += (
                                f"<div class='score-item'>"
                                f"<span class='score-label'>{score_name}:</span>"
                                f"{score_value}</div>"
                            )

                    html += "</div>"

                html += "</div>"

            html += "</div>"

        html += """
        </body>
        </html>
        """

        return html

    def _build_lesson_plan_lines(self, lesson_plan) -> list[str]:
        lines = [
            f"教案导出: {lesson_plan.title or ''}",
            "",
            f"学科: {lesson_plan.subject or ''}",
            f"年级: {lesson_plan.grade or ''}",
            f"课时时长: {lesson_plan.duration} 分钟",
            f"状态: {getattr(lesson_plan.status, 'value', lesson_plan.status)}",
            "",
            "教学目标:",
            lesson_plan.teaching_objectives or "无",
            "",
            "教学内容:",
            lesson_plan.teaching_content or "无",
            "",
            "教学方法:",
            lesson_plan.teaching_methods or "无",
            "",
            "教学过程:",
            lesson_plan.teaching_process or "无",
            "",
            "教学资源:",
            lesson_plan.teaching_resources or "无",
            "",
            "备注:",
            lesson_plan.notes or "无",
        ]
        return lines

    def _build_portfolio_report_lines(self, student, portfolios) -> list[str]:
        lines = [
            f"{student.name or ''} 的成长报告",
            "",
            "学生基本信息",
            f"姓名: {student.name or ''}",
            f"性别: {student.gender or ''}",
            f"出生日期: {student.birth_date}",
            f"年级: {student.grade or ''}",
            f"班级: {student.class_name or ''}",
            f"家长联系方式: {student.parent_contact or '未提供'}",
            "",
            "成长记录",
        ]

        type_names = {
            "work": "作品",
            "evaluation": "评价",
            "observation": "观察",
            "milestone": "里程碑",
        }

        for index, portfolio in enumerate(portfolios, start=1):
            lines.extend(
                [
                    "",
                    f"{index}. {portfolio.title or ''}",
                    f"类型: {type_names.get(portfolio.type, portfolio.type)}",
                    f"创建时间: {portfolio.created_at.strftime('%Y-%m-%d %H:%M:%S')}",
                    f"内容: {portfolio.content or '无'}",
                ]
            )

            scores = [
                ("认知", portfolio.cognitive_score),
                ("技能", portfolio.skill_score),
                ("创意", portfolio.creativity_score),
                ("合作", portfolio.cooperation_score),
                ("注意力", portfolio.attention_score),
            ]
            valid_scores = [f"{label}: {value}" for label, value in scores if value is not None]
            if valid_scores:
                lines.append("评分: " + " / ".join(valid_scores))

        if len(portfolios) == 0:
            lines.append("暂无成长记录")

        return lines

    def _render_simple_pdf(self, lines: list[str]) -> bytes:
        # 优先使用 reportlab（纯 Python、跨平台、内置中文 CID 字体）；
        # 仅当 reportlab 不可用时才退回手写极简 PDF。
        try:
            return self._render_simple_pdf_reportlab(lines)
        except Exception:
            return self._render_simple_pdf_legacy(lines)

    def _render_simple_pdf_reportlab(self, lines: list[str]) -> bytes:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.utils import simpleSplit
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfgen import canvas

        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))

        buffer = BytesIO()
        page_width, page_height = A4
        margin = 50
        font_size = 11
        line_height = 18
        c = canvas.Canvas(buffer, pagesize=A4)
        c.setTitle("个别化教育计划 IEP")

        usable_width = page_width - 2 * margin
        max_lines = int((page_height - 2 * margin) // line_height)

        def start_page():
            c.setFont("STSong-Light", font_size)
            return page_height - margin, 0

        y, count = start_page()
        for raw in lines:
            line = (raw or "").replace("\r", " ").replace("\n", " ")
            wrapped = simpleSplit(line, "STSong-Light", font_size, usable_width) or [""]
            for piece in wrapped:
                if count >= max_lines:
                    c.showPage()
                    y, count = start_page()
                c.drawString(margin, y, piece)
                y -= line_height
                count += 1
        c.showPage()
        c.save()
        return buffer.getvalue()

    def _render_simple_pdf_legacy(self, lines: list[str]) -> bytes:
        page_width = 595
        page_height = 842
        margin_left = 50
        start_y = 792
        line_height = 18
        max_lines_per_page = 38

        pages = [
            lines[i : i + max_lines_per_page] for i in range(0, len(lines), max_lines_per_page)
        ] or [[]]

        objects: list[bytes] = []

        def add_object(content: str) -> int:
            objects.append(content.encode("utf-8"))
            return len(objects)

        catalog_id = add_object("<< /Type /Catalog /Pages 2 0 R >>")
        pages_id = 2
        font_id = add_object(
            "<< /Type /Font /Subtype /Type0 /BaseFont /STSong-Light "
            "/Encoding /UniGB-UCS2-H /DescendantFonts [4 0 R] >>"
        )
        _ = add_object(
            "<< /Type /Font /Subtype /CIDFontType0 /BaseFont /STSong-Light "
            "/CIDSystemInfo << /Registry (Adobe) /Ordering (GB1) /Supplement 4 >> /DW 1000 >>"
        )

        page_object_ids: list[int] = []
        content_object_ids: list[int] = []

        for page_lines in pages:
            stream_lines = ["BT", "/F1 12 Tf", f"1 0 0 1 {margin_left} {start_y} Tm"]
            for index, line in enumerate(page_lines):
                safe_line = (line or "").replace("\r", " ").replace("\n", " ")
                hex_text = safe_line.encode("utf-16-be").hex().upper()
                if index > 0:
                    stream_lines.append(f"1 0 0 1 {margin_left} {start_y - line_height * index} Tm")
                stream_lines.append(f"<{hex_text}> Tj")
            stream_lines.append("ET")
            stream = "\n".join(stream_lines).encode("utf-8")
            content_id = add_object(
                f"<< /Length {len(stream)} >>\nstream\n{stream.decode('utf-8')}\nendstream"
            )
            content_object_ids.append(content_id)
            page_id = add_object(
                f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 {page_width} {page_height}] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
            )
            page_object_ids.append(page_id)

        kids = " ".join(f"{page_id} 0 R" for page_id in page_object_ids)
        objects[pages_id - 1] = (
            f"<< /Type /Pages /Kids [{kids}] /Count {len(page_object_ids)} >>".encode("utf-8")
        )

        pdf = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        offsets = [0]
        for index, obj in enumerate(objects, start=1):
            offsets.append(len(pdf))
            pdf.extend(f"{index} 0 obj\n".encode("utf-8"))
            pdf.extend(obj)
            pdf.extend(b"\nendobj\n")

        xref_start = len(pdf)
        pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode("utf-8"))
        pdf.extend(b"0000000000 65535 f \n")
        for offset in offsets[1:]:
            pdf.extend(f"{offset:010d} 00000 n \n".encode("utf-8"))

        pdf.extend(
            (
                f"trailer\n<< /Size {len(objects) + 1} /Root {catalog_id} 0 R >>\n"
                f"startxref\n{xref_start}\n%%EOF"
            ).encode("utf-8")
        )
        return bytes(pdf)

    # ==================================================
    # IEP 个别化教育计划导出
    # ==================================================

    _IEP_LEVEL_TEXT = {
        0: "不能完成",
        1: "大量辅助",
        2: "少量辅助",
        3: "独立完成",
    }
    _IEP_STATUS_TEXT = {"draft": "草稿", "review": "待审定", "locked": "已定稿"}
    _GOAL_STATUS_TEXT = {
        "not_started": "未开始",
        "in_progress": "进行中",
        "achieved": "已达成",
        "maintain": "维持",
    }

    def _iep_level(self, value: Any) -> str:
        if value is None:
            return "未评估"
        return f"{value} {self._IEP_LEVEL_TEXT.get(value, '')}".strip()

    def export_iep_to_pdf(self, iep: Any) -> bytes:
        """导出个别化教育计划（IEP）为 PDF；无 WeasyPrint 时回退内置简单 PDF。"""
        try:
            from weasyprint import CSS, HTML
        except (ImportError, OSError):
            return self._render_simple_pdf(self._build_iep_lines(iep))

        html_content = self._generate_iep_html(iep)
        css = CSS(string="""
            body {
                font-family: SimSun, SimHei, Arial, sans-serif;
                margin: 22px; color:#222; font-size: 12px;
            }
            h1 { text-align:center; font-size: 20px; margin-bottom: 4px; }
            .meta-sub { text-align:center; color:#666; margin-bottom: 18px; }
            h2 {
                font-size: 15px; border-bottom: 2px solid #548ca8;
                padding-bottom:4px; margin-top: 22px;
            }
            table { width:100%; border-collapse: collapse; margin: 8px 0 14px; }
            th, td {
                border:1px solid #bbb; padding:5px 7px; text-align:left;
                vertical-align: top;
            }
            th { background:#f0f4f7; }
            .domain-title { font-weight:bold; margin:12px 0 4px; }
            .overall { color:#548ca8; font-weight:bold; }
            .profile {
                background:#fafafa; border:1px solid #e5e5e5; padding:10px;
                border-radius:4px; white-space:pre-wrap;
            }
            .note { color:#777; }
        """)
        html_doc = HTML(string=html_content)
        pdf_bytes: bytes = html_doc.write_pdf(stylesheets=[css])  # type: ignore[no-any-return]
        return pdf_bytes

    def _iep_student_name(self, iep: Any) -> str:
        name = getattr(iep, "student_name", None)
        if name:
            return str(name)
        student = getattr(iep, "student", None)
        return str(getattr(student, "name", "") or "")

    def _generate_iep_html(self, iep: Any) -> str:
        name = escape(self._iep_student_name(iep))
        status = self._IEP_STATUS_TEXT.get(getattr(iep, "status", ""), "")
        reviewed = ""
        if getattr(iep, "reviewed_at", None):
            reviewed = getattr(iep.reviewed_at, "strftime", lambda _f: "")("%Y-%m-%d %H:%M")

        parts = [
            "<!DOCTYPE html><html><head><meta charset='UTF-8'>",
            f"<title>{name} 个别化教育计划</title></head><body>",
            "<h1>个别化教育计划（IEP）</h1>",
            f"<div class='meta-sub'>{name}　{escape(iep.term or '')}　"
            f"第 {iep.revision_no} 版　{escape(status)}</div>",
            "<table>",
            f"<tr><th style='width:120px'>学生</th><td>{name}</td>"
            f"<th style='width:120px'>学期</th><td>{escape(iep.term or '')}</td></tr>",
            f"<tr><th>状态</th><td>{escape(status)}</td>"
            f"<th>审定时间</th><td>{escape(reviewed) or '—'}</td></tr>",
            "</table>",
            "<h2>一、学生基本情况</h2>",
            f"<div class='profile'>{escape(iep.student_profile or '—')}</div>",
            "<h2>二、发展领域评估</h2>",
        ]

        if not iep.domain_assessments:
            parts.append("<p class='note'>暂无评估</p>")
        for d in iep.domain_assessments:
            parts.append(
                f"<div class='domain-title'>{escape(d.domain_name)}　"
                f"<span class='overall'>领域整体："
                f"{escape(self._iep_level(d.overall_level))}</span></div>"
            )
            rows = [
                "".join(
                    [
                        "<table><tr><th style='width:40%'>观察条目</th>"
                        "<th style='width:22%'>水平</th><th>备注</th></tr>",
                    ]
                )
            ]
            for it in d.items:
                rows.append(
                    "<tr>"
                    f"<td>{escape(it.item_name)}</td>"
                    f"<td>{escape(self._iep_level(it.level))}</td>"
                    f"<td>{escape(it.note or '') or '—'}</td>"
                    "</tr>"
                )
            rows.append("</table>")
            if d.summary:
                rows.append(f"<p class='note'>小结：{escape(d.summary)}</p>")
            parts.append("".join(rows))

        parts.append("<h2>三、长短期目标</h2>")
        if not iep.long_term_goals:
            parts.append("<p class='note'>暂无目标</p>")
        for g in iep.long_term_goals:
            parts.append(
                f"<div class='domain-title'>[{escape(g.domain_name or '未分领域')}] "
                f"长期目标：{escape(g.goal)}</div>"
            )
            table = [
                "<table><tr><th>条件</th><th>行为</th><th>标准</th>"
                "<th>支持级别</th><th>状态</th><th>教学方法</th></tr>"
            ]
            for s in g.short_term_goals:
                methods = "、".join(s.teaching_methods or [])
                table.append(
                    "<tr>"
                    f"<td>{escape(s.condition or '') or '—'}</td>"
                    f"<td>{escape(s.behavior)}</td>"
                    f"<td>{escape(s.criterion or '') or '—'}</td>"
                    f"<td>{escape(self._iep_level(s.support_level))}</td>"
                    f"<td>{escape(self._GOAL_STATUS_TEXT.get(s.status, s.status))}</td>"
                    f"<td>{escape(methods) or '—'}</td>"
                    "</tr>"
                )
            table.append("</table>")
            parts.append("".join(table))

        parts.append("<h2>四、支持措施</h2><table>")
        supports = iep.supports or {}
        for label, key in [("教学支持", "teaching"), ("康复支持", "rehab"), ("家庭支持", "family")]:
            value = "；".join(supports.get(key, []) or [])
            parts.append(
                f"<tr><th style='width:120px'>{label}</th><td>{escape(value) or '—'}</td></tr>"
            )
        parts.append("</table>")

        parts.append("<h2>五、评价方式</h2>")
        parts.append(f"<div class='profile'>{escape(iep.evaluation_plan or '—')}</div>")
        parts.append("</body></html>")
        return "".join(parts)

    def _build_iep_lines(self, iep: Any) -> list[str]:
        """无 WeasyPrint 时的纯文本回退。"""
        name = self._iep_student_name(iep)
        lines = [
            "个别化教育计划（IEP）",
            f"学生：{name}    学期：{iep.term}",
            f"版本：第 {iep.revision_no} 版    状态：{self._IEP_STATUS_TEXT.get(iep.status, '')}",
            "",
            "一、学生基本情况",
            iep.student_profile or "—",
            "",
            "二、发展领域评估",
        ]
        for d in iep.domain_assessments:
            lines.append(f"[{d.domain_name}] 整体：{self._iep_level(d.overall_level)}")
            for it in d.items:
                lines.append(f"  - {it.item_name}：{self._iep_level(it.level)}")
        lines.append("")
        lines.append("三、长短期目标")
        for g in iep.long_term_goals:
            lines.append(f"[{g.domain_name or '未分领域'}] 长期：{g.goal}")
            for s in g.short_term_goals:
                lines.append(
                    f"  · {s.condition or ''}{s.behavior}（{s.criterion or ''}）"
                    f" 支持：{self._iep_level(s.support_level)}"
                )
        lines.append("")
        lines.append("四、评价方式")
        lines.append(iep.evaluation_plan or "—")
        return lines
